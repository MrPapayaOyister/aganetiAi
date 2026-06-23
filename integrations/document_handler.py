import httpx
import json
import asyncio
import re
from pathlib import Path
from datetime import datetime, timezone

TEMP_DIR = Path("temp/uploads")
LLM_URL = "http://localhost:8080/v1/chat/completions"
SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md",
    ".doc", ".xls", ".ppt", ".rtf", ".odt", ".ods", ".odp"
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
MAX_CHARS = 12000  # max chars sent to LLM from extracted text

async def download_telegram_file(bot, file_id: str, filename: str, user_id: str) -> Path:
    """
    Downloads the file from Telegram and saves it to temp/uploads/{user_id}/ after sanitizing the filename.
    """
    user_dir = TEMP_DIR / user_id
    user_dir.mkdir(parents=True, exist_ok=True)

    # Sanitise filename: strip path separators, keep only [a-zA-Z0-9._-] characters
    clean_filename = re.sub(r'[^a-zA-Z0-9._-]', '', filename)
    if not clean_filename:
        clean_filename = "document"
        
    save_path = user_dir / clean_filename

    file_info = await bot.get_file(file_id)
    await bot.download_file(file_info.file_path, destination=save_path)
    return save_path

def extract_text(file_path: Path) -> str:
    """
    Extracts text from a file based on its extension, limiting the output to MAX_CHARS.
    """
    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
            return "\n\n".join(pages)[:MAX_CHARS]

        elif ext == ".docx":
            from docx import Document
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])[:MAX_CHARS]

        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                lines.append(f"[Sheet: {sheet}]")
                for row in ws.iter_rows(values_only=True):
                    row_str = " | ".join([str(c) if c is not None else "" for c in row])
                    if row_str.strip(" |"):
                        lines.append(row_str)
            return "\n".join(lines)[:MAX_CHARS]

        elif ext == ".csv":
            import pandas as pd
            df = pd.read_csv(file_path)
            return df.to_string(index=False)[:MAX_CHARS]

        elif ext in {".txt", ".md"}:
            return file_path.read_text(encoding="utf-8", errors="ignore")[:MAX_CHARS]

    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")
        return ""

    return ""

def analyse_document(text: str, filename: str, user_prompt: str = "") -> str:
    """
    Analyzes the document text using local LLM, asking custom prompt or default analysis prompt.
    """
    if user_prompt:
        system_content = "You are a document analyst. Answer the user's question about this document concisely and accurately."
        user_content = f"Document: {filename}\nContent:\n{text}\n\nQuestion: {user_prompt}"
    else:
        system_content = (
            "You are a document analyst. Analyse the document below and provide:\n"
            "1. A 2-sentence summary of what this document is about\n"
            "2. 3-5 key points or findings\n"
            "3. Any dates, deadlines, names, or numbers that stand out\n"
            "Be concise. Use bullet points."
        )
        user_content = f"Document: {filename}\nContent:\n{text}"

    payload = {
        "model": "local-model",
        "messages": [
            {"role": "system", "content": system_content},
            {"role": "user", "content": user_content}
        ],
        "max_tokens": 400,
        "temperature": 0.2
    }

    try:
        # Increased timeout to 90.0s since LLM can be slow
        response = httpx.post(LLM_URL, json=payload, timeout=90.0)
        if response.status_code == 200:
            return response.json()["choices"][0]["message"]["content"].strip()
    except Exception as e:
        print(f"Error calling LLM: {e}")
        
    return "⚠️ Could not analyse this document. Try asking a specific question about it."

async def analyse_document_stream(text: str, filename: str, user_prompt: str = ""):
    """
    Analyzes the document text using local LLM, yielding tokens dynamically.
    """
    if user_prompt:
        system_content = "You are a document analyst. Answer the user's question about this document concisely and accurately."
        user_content = f"Document: {filename}\nContent:\n{text}\n\nQuestion: {user_prompt}"
    else:
        system_content = (
            "You are a document analyst. Analyse the document below and provide:\n"
            "1. A 2-sentence summary of what this document is about\n"
            "2. 3-5 key points or findings\n"
            "3. Any dates, deadlines, names, or numbers that stand out\n"
            "Be concise. Use bullet points."
        )
        user_content = f"Document: {filename}\nContent:\n{text}"

    from integrations.model_router import route_model
    base_url = route_model(user_prompt or "analyse document")
    url = f"{base_url}/v1/chat/completions"

    payload = {
        "model": "local-model",
        "messages": [
            {"role": "system", "content": system_content},
            {"role": "user", "content": user_content}
        ],
        "max_tokens": 400,
        "temperature": 0.2,
        "stream": True
    }

    try:
        async with httpx.AsyncClient(timeout=90.0) as client:
            async with client.stream("POST", url, json=payload) as resp:
                if resp.status_code != 200:
                    yield f"Error: Backend returned status code {resp.status_code}."
                    return
                async for raw_line in resp.aiter_lines():
                    if not raw_line or raw_line.strip() == "data: [DONE]":
                        continue
                    try:
                        line = raw_line.removeprefix("data: ").strip()
                        data = json.loads(line)
                        token = data["choices"][0]["delta"].get("content", "")
                        if token:
                            yield token
                    except Exception:
                        continue
    except Exception as e:
        print(f"Error calling LLM in stream mode: {e}")
        yield "⚠️ Could not analyse this document. Try asking a specific question about it."

def build_followup_prompt(filename: str) -> str:
    """
    Returns the follow-up choices formatted string.
    """
    return (
        "What would you like to do with this document?\n\n"
        "• _\"summarise it\"_ — get a brief summary\n"
        "• _\"extract tasks\"_ — find action items and deadlines\n"
        "• _\"find all names\"_ — extract people mentioned\n"
        "• _\"draft a reply\"_ — write a response based on this document\n"
        "• _\"add to memory\"_ — save key facts to your long-term memory\n"
        "• Or ask any specific question about it"
    )

async def handle_document_upload(bot, message, user_id: str, status_msg=None) -> str:
    """
    Main entry point for document uploads in the Telegram bot.
    """
    if message.document:
        file_id = message.document.file_id
        filename = message.document.file_name or "document"
    elif message.photo:
        file_id = message.photo[-1].file_id
        filename = f"image_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
    else:
        return "⚠️ Unsupported file type."

    ext = Path(filename).suffix.lower()
    user_caption = message.caption or ""

    if ext in IMAGE_EXTENSIONS:
        return (
            "🖼️ Image received. Image analysis will be available with the "
            "Gemma 4 vision model (coming in Task 25). "
            "For now, try sending a PDF, DOCX, XLSX, CSV, or TXT file."
        )

    if ext not in SUPPORTED_EXTENSIONS:
        return (
            f"⚠️ File type `{ext}` is not supported yet.\n"
            f"Supported: PDF, DOCX, XLSX, CSV, TXT, MD, DOC, XLS, PPT, RTF, ODT, ODS, ODP"
        )

    # Download file
    try:
        file_path = await download_telegram_file(bot, file_id, filename, user_id)
    except Exception as e:
        return f"⚠️ Could not download file: {str(e)[:80]}"

    # Legacy format conversion using headless LibreOffice
    LEGACY_EXTENSIONS = {".doc", ".xls", ".ppt", ".rtf", ".odt", ".ods", ".odp"}
    if ext in LEGACY_EXTENSIONS:
        if status_msg:
            try:
                await status_msg.edit_text("🔄 Converting file format via LibreOffice...")
            except Exception:
                pass
        try:
            from integrations.libreoffice_converter import convert_to_pdf_safe
            converted_dir = file_path.parent / "converted"
            file_path = await convert_to_pdf_safe(file_path, converted_dir, timeout=25.0)
            filename = file_path.name
        except TimeoutError:
            return "⚠️ Conversion timed out. Please try uploading a smaller file."
        except Exception as e:
            return f"⚠️ Conversion failed: {str(e)[:100]}"

    # Extract text
    text = extract_text(file_path)
    if not text.strip():
        return (
            "⚠️ Could not extract text from this file. "
            "It may be scanned/image-based. Try a text-based PDF or DOCX."
        )

    # Store extracted text in session for follow-up questions
    session_file = TEMP_DIR / user_id / "last_document.json"
    session_file.parent.mkdir(parents=True, exist_ok=True)
    session_file.write_text(json.dumps({
        "filename": filename,
        "text": text,
        "uploaded_at": datetime.now(timezone.utc).isoformat()
    }, ensure_ascii=False))

    # Keep user engaged by updating status message right before calling slow LLM
    if status_msg:
        try:
            await status_msg.edit_text("🔍 Analysing your document...")
        except Exception:
            pass

    # Analyse using streaming
    collected = ""
    last_edit_at = 0.0
    token_count = 0

    async for token in analyse_document_stream(text, filename, user_caption):
        collected += token
        token_count += 1

        if status_msg and token_count % 10 == 0:
            import time
            now = time.time()
            if now - last_edit_at >= 1.0 and collected.strip():
                try:
                    display_text = collected.strip()
                    if "[ACTION:" in display_text:
                        display_text = display_text.split("[ACTION:")[0].strip()
                    if display_text:
                        await status_msg.edit_text(display_text + " ▌")
                        last_edit_at = now
                except Exception:
                    pass

    analysis = collected.strip()
    if "[ACTION:" in analysis:
        analysis = analysis.split("[ACTION:")[0].strip()
    analysis = analysis or "⚠️ No response received."

    if user_caption:
        return analysis
    else:
        return analysis + "\n\n" + build_followup_prompt(filename)
